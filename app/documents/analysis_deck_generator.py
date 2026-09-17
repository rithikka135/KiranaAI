from decimal import Decimal
from pathlib import Path

import matplotlib.pyplot as plt

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

from app.database.connection import SessionLocal
from app.services.analysis_service import get_weekly_analysis


OUTPUT_FOLDER = Path("generated_analysis")


def money(value):
    return f"₹{Decimal(str(value)):.2f}"


def add_title(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(
        Inches(0.6),
        Inches(0.3),
        Inches(12),
        Inches(0.6),
    )

    paragraph = title_box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True

    if subtitle:
        subtitle_box = slide.shapes.add_textbox(
            Inches(0.6),
            Inches(0.9),
            Inches(12),
            Inches(0.4),
        )

        paragraph = subtitle_box.text_frame.paragraphs[0]
        paragraph.text = subtitle
        paragraph.font.size = Pt(13)


def add_metric_card(
    slide,
    x,
    y,
    width,
    height,
    label,
    value,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )

    shape.text_frame.clear()

    p1 = shape.text_frame.paragraphs[0]
    p1.text = label
    p1.font.size = Pt(12)
    p1.font.bold = True
    p1.alignment = PP_ALIGN.CENTER

    p2 = shape.text_frame.add_paragraph()
    p2.text = value
    p2.font.size = Pt(20)
    p2.font.bold = True
    p2.alignment = PP_ALIGN.CENTER


def create_daily_sales_chart(data):
    dates = list(data["daily_sales"].keys())
    values = [
        float(Decimal(value))
        for value in data["daily_sales"].values()
    ]

    labels = [
        date_string[5:]
        for date_string in dates
    ]

    plt.figure(figsize=(10, 5))

    plt.plot(
        labels,
        values,
        marker="o",
    )

    plt.title("Daily Sales")
    plt.xlabel("Date")
    plt.ylabel("Sales (₹)")
    plt.xticks(rotation=45)
    plt.tight_layout()

    path = OUTPUT_FOLDER / "daily_sales.png"
    plt.savefig(path, dpi=150)
    plt.close()

    return path


def create_payment_chart(data):
    labels = [
        "Cash",
        "UPI",
        "Card",
        "Credit",
    ]

    values = [
        float(Decimal(data["cash_sales"])),
        float(Decimal(data["upi_sales"])),
        float(Decimal(data["card_sales"])),
        float(Decimal(data["credit_sales"])),
    ]

    plt.figure(figsize=(8, 5))

    plt.bar(
        labels,
        values,
    )

    plt.title("Payment Method Breakdown")
    plt.ylabel("Sales (₹)")
    plt.tight_layout()

    path = OUTPUT_FOLDER / "payment_breakdown.png"
    plt.savefig(path, dpi=150)
    plt.close()

    return path


def create_top_products_chart(data):
    products = data["top_products"][:5]

    if not products:
        return None

    names = [
        item["product_name"]
        for item in products
    ]

    revenues = [
        float(Decimal(item["revenue"]))
        for item in products
    ]

    names.reverse()
    revenues.reverse()

    plt.figure(figsize=(10, 5))

    plt.barh(
        names,
        revenues,
    )

    plt.title("Top 5 Products by Revenue")
    plt.xlabel("Revenue (₹)")
    plt.tight_layout()

    path = OUTPUT_FOLDER / "top_products.png"
    plt.savefig(path, dpi=150)
    plt.close()

    return path


def generate_analysis_deck(
    end_date=None,
) -> str:

    OUTPUT_FOLDER.mkdir(
        parents=True,
        exist_ok=True,
    )

    db = SessionLocal()

    try:

        analysis = get_weekly_analysis(
            db=db,
            end_date=end_date,
        )

    finally:

        db.close()

    daily_chart = create_daily_sales_chart(
        analysis
    )

    payment_chart = create_payment_chart(
        analysis
    )

    top_products_chart = create_top_products_chart(
        analysis
    )

    presentation = Presentation()

    # --------------------------------------------------
    # SLIDE 1 — TITLE
    # --------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )

    add_title(
        slide,
        "KiranaAI Weekly Sales Analysis",
        f"{analysis['start_date']} to {analysis['end_date']}",
    )

    textbox = slide.shapes.add_textbox(
        Inches(1),
        Inches(2),
        Inches(11),
        Inches(2),
    )

    paragraph = textbox.text_frame.paragraphs[0]
    paragraph.text = "Supermarket Operations Report"
    paragraph.font.size = Pt(30)
    paragraph.font.bold = True
    paragraph.alignment = PP_ALIGN.CENTER

    # --------------------------------------------------
    # SLIDE 2 — SALES OVERVIEW
    # --------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )

    add_title(
        slide,
        "Sales Overview",
    )

    add_metric_card(
        slide,
        0.7,
        1.5,
        2.8,
        1.5,
        "Total Sales",
        money(analysis["total_sales"]),
    )

    add_metric_card(
        slide,
        3.7,
        1.5,
        2.8,
        1.5,
        "Bills",
        str(analysis["bill_count"]),
    )

    add_metric_card(
        slide,
        6.7,
        1.5,
        2.8,
        1.5,
        "GST Collected",
        money(analysis["total_gst"]),
    )

    add_metric_card(
        slide,
        9.7,
        1.5,
        2.8,
        1.5,
        "Products",
        str(analysis["total_products"]),
    )

    slide.shapes.add_picture(
        str(daily_chart),
        Inches(1),
        Inches(3.4),
        width=Inches(11),
    )

    # --------------------------------------------------
    # SLIDE 3 — PAYMENT BREAKDOWN
    # --------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )

    add_title(
        slide,
        "Payment Method Breakdown",
    )

    slide.shapes.add_picture(
        str(payment_chart),
        Inches(2),
        Inches(1.5),
        width=Inches(9),
    )

    # --------------------------------------------------
    # SLIDE 4 — TOP PRODUCTS
    # --------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )

    add_title(
        slide,
        "Top Selling Products",
    )

    if top_products_chart:

        slide.shapes.add_picture(
            str(top_products_chart),
            Inches(1),
            Inches(1.4),
            width=Inches(11),
        )

    else:

        textbox = slide.shapes.add_textbox(
            Inches(2),
            Inches(3),
            Inches(8),
            Inches(1),
        )

        paragraph = textbox.text_frame.paragraphs[0]
        paragraph.text = "No sales data available."
        paragraph.font.size = Pt(20)
        paragraph.alignment = PP_ALIGN.CENTER

    # --------------------------------------------------
    # SLIDE 5 — INVENTORY HEALTH
    # --------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )

    add_title(
        slide,
        "Inventory Health",
    )

    add_metric_card(
        slide,
        1,
        1.5,
        3.5,
        1.5,
        "Active Products",
        str(analysis["total_products"]),
    )

    add_metric_card(
        slide,
        5,
        1.5,
        3.5,
        1.5,
        "Low Stock Items",
        str(
            len(
                analysis["low_stock_products"]
            )
        ),
    )

    textbox = slide.shapes.add_textbox(
        Inches(1),
        Inches(3.5),
        Inches(11),
        Inches(3),
    )

    frame = textbox.text_frame
    frame.clear()

    if analysis["low_stock_products"]:

        p = frame.paragraphs[0]
        p.text = "Products requiring attention:"
        p.font.size = Pt(18)
        p.font.bold = True

        for item in analysis["low_stock_products"]:

            p = frame.add_paragraph()

            p.text = (
                f"{item['product_name']} — "
                f"{item['quantity']} {item['unit']} "
                f"(reorder level: "
                f"{item['reorder_level']})"
            )

            p.font.size = Pt(14)

    else:

        p = frame.paragraphs[0]
        p.text = "No products are currently below reorder level."
        p.font.size = Pt(18)

    # --------------------------------------------------
    # SLIDE 6 — GST
    # --------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )

    add_title(
        slide,
        "GST Collection",
    )

    add_metric_card(
        slide,
        1,
        1.5,
        3.5,
        1.5,
        "CGST",
        money(analysis["cgst"]),
    )

    add_metric_card(
        slide,
        5,
        1.5,
        3.5,
        1.5,
        "SGST",
        money(analysis["sgst"]),
    )

    add_metric_card(
        slide,
        9,
        1.5,
        3,
        1.5,
        "Total GST",
        money(analysis["total_gst"]),
    )

    # --------------------------------------------------
    # SLIDE 7 — INSIGHTS
    # --------------------------------------------------

    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )

    add_title(
        slide,
        "Key Business Insights",
    )

    textbox = slide.shapes.add_textbox(
        Inches(1),
        Inches(1.5),
        Inches(11),
        Inches(5),
    )

    frame = textbox.text_frame
    frame.clear()

    for index, insight in enumerate(
        analysis["insights"]
    ):

        if index == 0:
            paragraph = frame.paragraphs[0]
        else:
            paragraph = frame.add_paragraph()

        paragraph.text = f"• {insight}"
        paragraph.font.size = Pt(18)
        paragraph.space_after = Pt(15)

    # --------------------------------------------------
    # SAVE
    # --------------------------------------------------

    output_path = (
        OUTPUT_FOLDER
        / (
            f"weekly_sales_analysis_"
            f"{analysis['start_date']}_"
            f"{analysis['end_date']}.pptx"
        )
    )

    presentation.save(
        output_path
    )

    return str(output_path)